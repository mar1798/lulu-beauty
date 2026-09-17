"""Full catalogue → xlsx, the mirror image of the xlsx import.

The column names are the import's own (`name`, `slug`, `price`, …) rather than Russian
captions, because the point of this sheet is the round trip: the owner exports what the
shop has, edits prices or stock in Excel, and uploads the same file back on the page the
export button lives on. A Russian header would break that on the way back in.

`description` is deliberately absent. It is the one field that is paragraphs rather than
a value, it turns every row into something no longer readable on screen, and nobody edits
it in a spreadsheet — it is written on the product page in the admin panel. Leaving the
column out is safe for the round trip: the import treats an absent column as "leave it
alone", so re-uploading this file keeps the descriptions the catalogue already has.

Photos are not in the sheet either. xlsx can only carry them as floating pictures anchored
to a cell — which Google Sheets drops on import, and which cost a disk read and a re-encode
per product — and they are no help in what the sheet is for: editing values and uploading
the file back.
"""

import functools
import io
from dataclasses import dataclass
from datetime import datetime
from typing import cast
from zoneinfo import ZoneInfo

import anyio.to_thread
from openpyxl import Workbook
from openpyxl.cell import Cell
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.catalog.models import Product
from app.config import settings
from app.export.service import (
    MAX_TEXT_WIDTH,
    MONEY_FORMAT,
    WIDTH_PADDING,
    CellValue,
)

SHEET_TITLE = "Каталог"

HEADER = ["name", "slug", "brand", "category", "price", "volume", "inStock"]

# Column indexes, by the header above. Named because both width and alignment are keyed
# off them.
TEXT_COLUMNS = (0, 1, 2, 3)
PRICE_COLUMN = 4
VOLUME_COLUMN = 5

QUANTITY_FORMAT = "#,##0"

# What the import reads back as a boolean (`TRUTHY_VALUES`/`FALSY_VALUES`). Written as
# words rather than TRUE/FALSE because the sheet is read by a person first.
IN_STOCK_YES = "да"
IN_STOCK_NO = "нет"


@dataclass(frozen=True)
class ProductExportRow:
    """One catalogue line, in the shape the import reads back."""

    name: str
    slug: str
    brand: str | None
    category_slug: str | None
    price_cents: int
    volume_ml: int | None
    in_stock: bool


def export_filename() -> str:
    """`catalog-2030-06-01.xlsx`, dated in the shop's own calendar.

    Dated because the owner downloads this repeatedly and edits it offline, and a folder
    of `catalog.xlsx`, `catalog (1).xlsx`, `catalog (2).xlsx` says nothing about which one
    is the file they were about to upload back. The shop's timezone, not the server's: the
    date on the file has to be the date the owner would write on it.
    """
    today = datetime.now(ZoneInfo(settings.cycle_timezone)).date()
    return f"catalog-{today.isoformat()}.xlsx"


def _cell_values(row: ProductExportRow) -> list[CellValue]:
    return [
        row.name,
        row.slug,
        row.brand or "",
        row.category_slug or "",
        row.price_cents / 100,
        row.volume_ml,
        IN_STOCK_YES if row.in_stock else IN_STOCK_NO,
    ]


def _formats() -> list[str | None]:
    formats: list[str | None] = [None] * len(HEADER)
    formats[PRICE_COLUMN] = MONEY_FORMAT
    formats[VOLUME_COLUMN] = QUANTITY_FORMAT
    return formats


def _rendered_width(value: CellValue, number_format: str | None) -> int:
    if isinstance(value, float) and number_format == MONEY_FORMAT:
        return len(f"{value:,.2f}")
    if value is None:
        return 0
    return len(str(value))


def _apply_widths(sheet: Worksheet, body: list[list[CellValue]]) -> None:
    formats = _formats()
    widths = [len(title) for title in HEADER]
    for values in body:
        for index, value in enumerate(values):
            widths[index] = max(widths[index], _rendered_width(value, formats[index]))

    for index, width in enumerate(widths):
        limit = MAX_TEXT_WIDTH if index in TEXT_COLUMNS else width
        sheet.column_dimensions[get_column_letter(index + 1)].width = min(width, limit) + (
            WIDTH_PADDING
        )


def _write_cell(
    sheet: Worksheet, line: int, value: CellValue, *, index: int, bold: bool = False
) -> None:
    # `Worksheet.cell` is typed as possibly returning a `MergedCell`, which has no
    # writable value. This sheet is built from nothing and merges no cells.
    cell = cast(Cell, sheet.cell(row=line, column=index + 1))
    # Same trap as the orders sheet: openpyxl writes a string starting with "=" as a real
    # formula, and a product name comes from a supplier's price list. See the note in
    # `export/service.py` — declaring the type keeps the text exactly as typed.
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"
    cell.alignment = Alignment(
        wrap_text=index in TEXT_COLUMNS,
        vertical="center",
        horizontal="left" if index in TEXT_COLUMNS else "right",
    )
    number_format = _formats()[index]
    if number_format is not None:
        cell.number_format = number_format
    if bold:
        cell.font = Font(bold=True)


def build_products_workbook(rows: list[ProductExportRow]) -> bytes:
    """Serializes the catalogue to xlsx bytes. Blocking CPU — call it off the event loop.

    Not `write_only`, unlike the orders sheet. What buys write-only mode its speed is
    streaming each row straight into the archive, and it costs the ability to touch a row
    after it is written; the trade is worth it for an all-cycles order export of hundreds
    of thousands of lines, and pointless for a catalogue of a few thousand.
    """
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = SHEET_TITLE

    for index, title in enumerate(HEADER):
        _write_cell(sheet, 1, title, index=index, bold=True)

    body = [_cell_values(row) for row in rows]
    for offset, values in enumerate(body):
        for index, value in enumerate(values):
            _write_cell(sheet, offset + 2, value, index=index)

    _apply_widths(sheet, body)
    # The header stays visible while the owner scrolls a catalogue of any length — the
    # columns are `slug` and `inStock`, which are not guessable from the values.
    sheet.freeze_panes = "A2"

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class CatalogExportService:
    def __init__(self, session: AsyncSession) -> None:
        self._session = session

    async def export_products(self) -> tuple[bytes, str]:
        rows = await self._rows()
        # In a worker thread, for the reason the orders export gives: the serialization is
        # pure CPU, and the API process is single-threaded — a blocking second here is a
        # second of downtime for the whole shop, not just for the owner downloading.
        content = await anyio.to_thread.run_sync(functools.partial(build_products_workbook, rows))
        return content, export_filename()

    async def _rows(self) -> list[ProductExportRow]:
        """Every live product, grouped by brand — the order the owner edits a price list in.

        Not the admin catalogue's order (`name`, `id`): a sheet is worked through one
        supplier at a time, and a screen is scrolled looking for one product by name.

        `id` last, and not for show. Product names are not unique — two volumes of the
        same toner, a re-imported duplicate — so `(brand, name)` leaves the order of the
        equal rows to the planner, and two downloads of an unchanged catalogue then differ
        in row order. This file is diffed against the previous one.

        Soft-deleted products are left out: the sheet is meant to be edited and uploaded
        back, and a deleted row coming back through the import would quietly resurrect the
        product (the import matches on slug and knows nothing about `deleted_at`).
        """
        query = (
            select(Product)
            .where(Product.deleted_at.is_(None))
            .options(selectinload(Product.category))
            .order_by(Product.brand.nulls_last(), Product.name, Product.id)
        )
        products = (await self._session.execute(query)).scalars().all()

        return [
            ProductExportRow(
                name=product.name,
                slug=product.slug,
                brand=product.brand,
                # The slug, not the name: the import reads either, but only the slug is
                # the identity — two categories may read the same to a person.
                category_slug=product.category.slug if product.category else None,
                price_cents=product.price_cents,
                volume_ml=product.volume_ml,
                in_stock=product.in_stock,
            )
            for product in products
        ]
