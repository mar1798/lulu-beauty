import io
from datetime import date, datetime
from zoneinfo import ZoneInfo

import openpyxl

from app.catalog.import_service import (
    parse_in_stock,
    parse_price_cents,
    parse_rows,
    parse_volume_ml,
    validate_headers,
)
from app.catalog.shades import ShadeLabel, normalize_shade, parse_shade_label
from app.config import settings
from app.export.products import (
    HEADER,
    ProductExportRow,
    build_products_workbook,
    export_filename,
)


def _row(**overrides: object) -> ProductExportRow:
    defaults: dict[str, object] = {
        "name": "Rose Serum",
        "slug": "rose-serum",
        "brand": "Lulu",
        "category_slug": "care",
        "price_cents": 15000,
        "volume_ml": 50,
        "shade": None,
        "shade_label": ShadeLabel.TONE,
        "in_stock": True,
    }
    defaults.update(overrides)
    return ProductExportRow(**defaults)  # type: ignore[arg-type]


def _sheet(rows: list[ProductExportRow]):  # type: ignore[no-untyped-def]
    workbook = openpyxl.load_workbook(io.BytesIO(build_products_workbook(rows)))
    sheet = workbook.active
    assert sheet is not None
    return sheet


def _values(sheet, index: int) -> list[object]:  # type: ignore[no-untyped-def]
    return [cell.value for cell in next(sheet.iter_rows(min_row=index, max_row=index))]


def test_header_is_the_import_s_own_column_names() -> None:
    """The sheet is meant to be edited and uploaded back, so the header has to be readable
    by the import — a Russian caption would break the round trip."""
    assert _values(_sheet([]), 1) == HEADER


def test_description_and_photos_are_not_columns() -> None:
    """Both are edited on the product page, not in a spreadsheet — and the import leaves
    a column it does not see alone, so neither is lost by uploading this file back."""
    assert "description" not in HEADER
    assert "image" not in HEADER


def test_writes_a_line_per_product() -> None:
    sheet = _sheet([_row(name="Крем", price_cents=15000, volume_ml=50)])

    assert _values(sheet, 2) == ["Крем", "rose-serum", "Lulu", "care", 150, 50, None, None, "да"]


def test_writes_the_shade_of_a_shaded_variant() -> None:
    sheet = _sheet([_row(name="Помада", volume_ml=None, shade="05 Rose")])

    assert _values(sheet, 2)[HEADER.index("shade")] == "05 Rose"


def test_writes_what_the_shades_are_called_in_russian() -> None:
    """The word, not the code: the sheet is read by a person, and the import takes it back."""
    sheet = _sheet([_row(shade="Red", shade_label=ShadeLabel.COLOR)])

    assert _values(sheet, 2)[HEADER.index("shadeLabel")] == "Цвет"


def test_leaves_the_label_empty_on_a_product_without_shades() -> None:
    """Nothing on the page says it there — and an empty cell leaves it alone on the way
    back, so the owner's choice survives the round trip all the same."""
    sheet = _sheet([_row(shade=None, shade_label=ShadeLabel.SCENT)])

    assert _values(sheet, 2)[HEADER.index("shadeLabel")] is None


def test_missing_brand_volume_and_category_come_out_empty() -> None:
    sheet = _sheet([_row(brand=None, category_slug=None, volume_ml=None, in_stock=False)])
    _, _, brand, category, _, volume, shade, label, in_stock = _values(sheet, 2)

    assert (brand, category, volume, shade, label) == (None, None, None, None, None)
    assert in_stock == "нет"


def test_export_round_trips_through_the_import() -> None:
    rows = [
        _row(name="Крем", slug="krem", price_cents=15000, volume_ml=50, in_stock=True),
        _row(name="Тоник", slug="tonik", brand=None, category_slug=None, volume_ml=None,
             price_cents=999, in_stock=False),
        _row(name="Помада", slug="pomada", volume_ml=None, shade="05 Rose", price_cents=700,
             shade_label=ShadeLabel.SHADE),
    ]  # fmt: skip

    parsed = parse_rows("catalog.xlsx", build_products_workbook(rows))

    assert validate_headers(parsed) is None
    assert [
        (
            cells["name"],
            cells["slug"],
            parse_price_cents(cells["price"]),
            parse_volume_ml(cells["volume"]),
            normalize_shade(cells["shade"]),
            parse_shade_label(cells["shade_label"]),
            parse_in_stock(cells["in_stock"]),
        )
        for _, cells in parsed
    ] == [
        ("Крем", "krem", 15000, 50, None, None, True),
        ("Тоник", "tonik", 999, None, None, None, False),
        ("Помада", "pomada", 700, None, "05 Rose", ShadeLabel.SHADE, True),
    ]


def test_a_product_name_that_looks_like_a_formula_stays_text() -> None:
    sheet = _sheet([_row(name="=1+1")])
    cell = sheet.cell(row=2, column=HEADER.index("name") + 1)

    assert cell.data_type == "s"
    assert cell.value == "=1+1"


def test_header_stays_visible_while_the_catalogue_is_scrolled() -> None:
    assert _sheet([]).freeze_panes == "A2"


def test_export_filename_is_dated_so_the_owner_can_tell_two_downloads_apart() -> None:
    first, _, rest = export_filename().partition("catalog-")

    assert first == ""
    assert rest.endswith(".xlsx")
    assert (
        date.fromisoformat(rest.removesuffix(".xlsx"))
        == datetime.now(ZoneInfo(settings.cycle_timezone)).date()
    )
